import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from pa_workflow import p1db,at_pout,execute,validate,sim_dc,sim_rf

ROOT=Path(__file__).resolve().parents[1]

def config():return json.loads((ROOT/'workflow.example.json').read_text(encoding='utf-8'))

class WorkflowTests(unittest.TestCase):
    def test_p1db_interpolation(self):
        rows=[dict(pin_dbm=x,gain_db=g) for x,g in [(0,20),(1,20),(2,20),(3,19.5),(4,18.5)]]
        r=p1db(rows)
        self.assertEqual(r['p1db_input_dbm'],3.5)
        self.assertEqual(r['p1db_output_dbm'],22.5)
        rows[-1]['gain_db']=19.7
        self.assertEqual(p1db(rows)['status'],'not_bracketed')
    def test_bad_baseline_and_duplicates(self):
        self.assertEqual(p1db([dict(pin_dbm=x,gain_db=20-x) for x in range(5)])['status'],'baseline_not_flat')
        self.assertEqual(p1db([dict(pin_dbm=0,gain_db=20)]*5)['status'],'duplicate_pin')
    def test_pout_no_extrapolation(self):
        rows=[dict(pin_dbm=0,pout_dbm=30,total_idd_a=1,aclr_lower_dbc=-50,aclr_upper_dbc=-49),dict(pin_dbm=1,pout_dbm=32,total_idd_a=2,aclr_lower_dbc=-40,aclr_upper_dbc=-39)]
        r,status=at_pout(rows,31)
        self.assertEqual(r['total_idd_a'],1.5)
        self.assertEqual(r['aclr_upper_dbc'],-44)
        self.assertEqual(at_pout(rows,40)[1],'target_not_bracketed')
        rows[1]['pout_dbm']=29
        self.assertEqual(at_pout(rows,30)[1],'nonmonotonic_pout')
    def test_excel_records_charts_and_rank(self):
        with tempfile.TemporaryDirectory() as d:
            path=execute(config(),Path(d)/'run')
            wb=load_workbook(path)
            self.assertEqual(wb['靜態資料'].max_row,10)
            self.assertEqual(wb['RF資料'].max_row,370)
            self.assertEqual(len(wb['曲線圖']._charts),7)
            self.assertEqual(len(wb['曲線圖']._charts[0].series),9)
            self.assertGreater(wb['電流事件'].max_row,1)
            statuses=[r[12] for r in wb['工作點比較'].iter_rows(min_row=2,values_only=True)]
            self.assertIn('qualified',statuses)
            wb.close()
    def test_missing_hardware_prevents_execution(self):
        c=config();c['mode']='hardware'
        with tempfile.TemporaryDirectory() as d:
            target=Path(d)/'run'
            with self.assertRaisesRegex(ValueError,'替換實機命令'):
                execute(c,target)
            self.assertFalse(target.exists())
    def test_missing_criteria_not_recommended(self):
        c=config();c['analysis']['aclr_limit_dbc']=None
        with tempfile.TemporaryDirectory() as d:
            wb=load_workbook(execute(c,Path(d)/'run'))
            for r in wb['工作點比較'].iter_rows(min_row=2,values_only=True):
                self.assertNotEqual(r[12],'qualified');self.assertIsNone(r[13])
            wb.close()
    def test_hardware_failure_cleanup_partial_data(self):
        c=config();c['mode']='hardware'
        c['commands']={k:['mock'] for k in c['commands']}
        c['commands']['set_bias']=['mock','{main}','{peaking}'];c['commands']['set_rf']=['mock','{power}']
        calls=[]
        def call(self,key,**values):
            calls.append(key)
            if key=='read_dc':return sim_dc(c,c['bias_points'][0])
            if key=='read_rf':raise RuntimeError('test sensor failure')
        with tempfile.TemporaryDirectory() as d,patch('pa_workflow.CommandAdapter.call',call),patch('pa_workflow.time.sleep'):
            target=Path(d)/'run'
            with self.assertRaisesRegex(RuntimeError,'部分資料'):
                execute(c,target)
            self.assertEqual(calls[-2:],['rf_off','safe_bias'])
            self.assertLess(calls.index('set_rf'),calls.index('rf_on'))
            self.assertTrue((target/'report.xlsx').exists())
            status=json.loads((target/'run_status.json').read_text())
            self.assertFalse(status['completed'])

class ConfigSaveTests(unittest.TestCase):
    def test_save_and_update_without_hardware(self):
        from pa_workflow import save_workflow_config
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'settings.json'
            c=config();c['mode']='hardware';c['waveform']='測試波形'
            # Incomplete hardware commands are valid drafts; no commands run.
            save_workflow_config(json.dumps(c,ensure_ascii=False),p)
            self.assertEqual(json.loads(p.read_text(encoding='utf-8')),c)
            c['analysis']['target_pout_dbm']=46
            save_workflow_config(json.dumps(c),p)
            self.assertEqual(json.loads(p.read_text())['analysis']['target_pout_dbm'],46)
    def test_invalid_json_preserves_existing_file(self):
        from pa_workflow import save_workflow_config
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'settings.json';p.write_text('{"original":true}')
            for value in ('{broken','[]','{"value":NaN}'):
                with self.assertRaises(ValueError):
                    save_workflow_config(value,p)
                self.assertEqual(p.read_text(),'{"original":true}')
            self.assertEqual(list(Path(d).iterdir()),[p])
