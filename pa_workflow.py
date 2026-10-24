"""Doherty characterization: configurable command adapter, analysis and Excel charts."""
import argparse
import csv
import json
import math
import subprocess
import time
from datetime import datetime, timezone
from pathlib import Path

FIELDS = ['mode','stage','timestamp_utc','bias_id','main_vgs_set_v','peaking_vgs_set_v',
          'main_vgs_v','peaking_vgs_v','main_vds_v','peaking_vds_v','main_idd_a',
          'peaking_idd_a','total_idd_a','temperature_c','rf_set_dbm','pin_dbm','pout_dbm',
          'gain_db','aclr_lower_dbc','aclr_upper_dbc','event','frequency_hz','waveform','load_description']


def finite(value):
    x=float(value)
    if not math.isfinite(x):
        raise ValueError('量測／設定必須為有限數值')
    return x


def axis(start,stop,step):
    from decimal import Decimal
    a,b,s=map(lambda x:Decimal(str(x)),(start,stop,step))
    if not all(v.is_finite() for v in (a,b,s)) or s<=0:
        raise ValueError('掃描步距須為有限正數')
    n=int(abs(b-a)//s)+1
    if n>10000:
        raise ValueError('單軸點數過多')
    return [float(a+i*s*(1 if b>=a else -1)) for i in range(n)]


def validate(c):
    if c['mode'] not in ('simulation','hardware'):
        raise ValueError('mode 必須是 simulation 或 hardware')
    for name in ('main','peaking'):
        lo,hi=map(finite,c['limits'][name])
        if lo>hi:
            raise ValueError('RD 電壓上下限順序錯誤')
        for point in c['bias_points']:
            if not lo<=finite(point[name])<=hi:
                raise ValueError(name+' 偏壓超出 RD 範圍')
    if not c['bias_points']:
        raise ValueError('沒有候選偏壓')
    ids=[(finite(p['main']),finite(p['peaking'])) for p in c['bias_points']]
    if len(set(ids))!=len(ids):
        raise ValueError('偏壓組合重複')
    powers=axis(**c['rf_sweep'])
    if min(powers)<0 or max(powers)>50:
        raise ValueError('RF 設定範圍須在 0～50 dBm')
    if len(powers)*len(ids)>100000:
        raise ValueError('最多 100000 個 RF 點')
    if finite(c['settle_seconds'])<0 or finite(c['command_timeout_seconds'])<=0:
        raise ValueError('等待時間或 timeout 無效')
    for key in ('frequency_hz','vdd_v'):
        if finite(c[key])<=0:
            raise ValueError(key+' 須 >0')
    if not c['waveform'] or not c['load_description']:
        raise ValueError('請填寫波形與負載條件')
    a=c['analysis']
    finite(a['target_pout_dbm'])
    for key in ('min_p1db_output_dbm','aclr_limit_dbc'):
        if a.get(key) is not None:
            finite(a[key])
    if int(a.get('baseline_points',3))<2:
        raise ValueError('baseline_points 至少 2')
    if finite(a.get('baseline_gain_tolerance_db',.2))<0:
        raise ValueError('小訊號增益容差不得為負')
    if c['mode']=='simulation':
        for p in c['bias_points']:
            if not (-5<=p['main']<=-1 and -10<=p['peaking']<=-1):
                raise ValueError('模擬偏壓僅支援 Main -5～-1 V、Peaking -10～-1 V')
    if c['mode']=='hardware':
        for key in ('set_bias','read_dc','rf_on','rf_off','set_rf','read_rf','safe_bias'):
            cmd=c.get('commands',{}).get(key)
            if not isinstance(cmd,list) or not cmd or not all(isinstance(v,str) and v for v in cmd):
                raise ValueError('硬體命令待設定：'+key+'（必須是 argv 字串陣列）')
            if any('REPLACE_' in v for v in cmd):
                raise ValueError('請先替換實機命令：'+key)
        if not all('{'+x+'}' in ' '.join(c['commands']['set_bias']) for x in ('main','peaking')):
            raise ValueError('set_bias 需包含 {main} 與 {peaking}')
        if '{power}' not in ' '.join(c['commands']['set_rf']):
            raise ValueError('set_rf 需包含 {power}')
    return powers


class CommandAdapter:
    def __init__(self,c):
        self.c=c
    def call(self,key,**values):
        argv=[part.format(**values) for part in self.c['commands'][key]]
        r=subprocess.run(argv,shell=False,capture_output=True,text=True,
                         timeout=self.c['command_timeout_seconds'])
        if r.returncode:
            raise RuntimeError(f'命令 {key} 執行失敗，exit={r.returncode}；請在本機檢查命令輸出。')
        if key.startswith('read_'):
            data=json.loads(r.stdout)
            if not isinstance(data,dict):
                raise ValueError(key+' 必須回傳 JSON object')
            return data


def dc_values(data):
    keys=['main_vgs_v','peaking_vgs_v','main_vds_v','peaking_vds_v',
          'main_idd_a','peaking_idd_a','temperature_c']
    out={k:finite(data[k]) for k in keys}
    if min(out['main_idd_a'],out['peaking_idd_a'])<0:
        raise ValueError('支路 Idd 必須為非負 A 值')
    out['total_idd_a']=out['main_idd_a']+out['peaking_idd_a']
    return out


def sim_dc(c,p,power=None):
    # Arbitrary, bounded demonstration. No device calibration.
    main=1.6/(1+math.exp(-(p['main']+3.66566)/.45))
    peak=.8/(1+math.exp(-(p['peaking']+4.0982)/.6))
    if power is not None:
        drive=1/(1+math.exp(-(power-25)/4))
        main+=1.2*drive;peak+=1.0*drive
    return dict(main_vgs_v=p['main'],peaking_vgs_v=p['peaking'],main_vds_v=c['vdd_v'],
                peaking_vds_v=c['vdd_v'],main_idd_a=main,peaking_idd_a=peak,temperature_c=35)


def sim_rf(p,power):
    g0=17+(p['main']+3)*1.5
    pin1=26+(p['main']+3)*2
    compression=10*math.log10(1+(10**.1-1)*10**((power-pin1)/10))
    pout=power+g0-compression
    aclr=-48+12/(1+math.exp(-(power-pin1)/3))
    return dict(pin_dbm=power,pout_dbm=pout,aclr_lower_dbc=aclr,aclr_upper_dbc=aclr+.5)


def capture(c,stage,dc,bias,index,power=None,rf=None):
    r={k:None for k in FIELDS}
    r.update(dc_values(dc))
    r.update(mode=c['mode'],stage=stage,timestamp_utc=datetime.now(timezone.utc).isoformat(),
             bias_id=index,main_vgs_set_v=bias['main'],peaking_vgs_set_v=bias['peaking'],
             rf_set_dbm=power,frequency_hz=c['frequency_hz'],waveform=c['waveform'],load_description=c['load_description'])
    if rf is not None:
        r.update({k:finite(rf[k]) for k in ['pin_dbm','pout_dbm','aclr_lower_dbc','aclr_upper_dbc']})
        r['gain_db']=r['pout_dbm']-r['pin_dbm']
    r['event']='total_current_ge_3A' if r['total_idd_a']>=3 else ''
    return r


def p1db(rows,baseline_points=3,tolerance=.2):
    ordered=sorted(rows,key=lambda r:r['pin_dbm'])
    if len(ordered)<baseline_points+1:
        return dict(status='insufficient_points')
    if len(set(r['pin_dbm'] for r in ordered))!=len(ordered):
        return dict(status='duplicate_pin')
    gains=[r['gain_db'] for r in ordered[:baseline_points]]
    if max(gains)-min(gains)>tolerance:
        return dict(status='baseline_not_flat')
    g0=sum(gains)/len(gains)
    for l,r in zip(ordered,ordered[1:]):
        cl,cr=g0-l['gain_db'],g0-r['gain_db']
        if cl<1<=cr:
            t=(1-cl)/(cr-cl)
            pin=l['pin_dbm']+t*(r['pin_dbm']-l['pin_dbm'])
            return dict(status='found',small_signal_gain_db=g0,p1db_input_dbm=pin,p1db_output_dbm=pin+g0-1)
    return dict(status='not_bracketed',small_signal_gain_db=g0)


def at_pout(rows,target):
    # No extrapolation; ambiguous/nonmonotonic output curves must be reviewed.
    ordered=sorted(rows,key=lambda r:r['pin_dbm'])
    if any(r['pout_dbm']<=l['pout_dbm'] for l,r in zip(ordered,ordered[1:])):
        return None,'nonmonotonic_pout'
    if not ordered or not ordered[0]['pout_dbm']<=target<=ordered[-1]['pout_dbm']:
        return None,'target_not_bracketed'
    for r in ordered:
        if r['pout_dbm']==target:
            return {k:r[k] for k in ('total_idd_a','aclr_lower_dbc','aclr_upper_dbc')},'exact'
    for l,r in zip(ordered,ordered[1:]):
        if l['pout_dbm']<target<r['pout_dbm']:
            t=(target-l['pout_dbm'])/(r['pout_dbm']-l['pout_dbm'])
            return {k:l[k]+t*(r[k]-l[k]) for k in ('total_idd_a','aclr_lower_dbc','aclr_upper_dbc')},'interpolated'
    return None,'target_not_bracketed'


def analyze(rows,c):
    a=c['analysis'];results=[]
    for index in sorted({r['bias_id'] for r in rows if r['stage']=='rf'}):
        group=[r for r in rows if r['stage']=='rf' and r['bias_id']==index]
        result=dict(bias_id=index,main_vgs_v=group[0]['main_vgs_set_v'],peaking_vgs_v=group[0]['peaking_vgs_set_v'])
        result.update(p1db(group,int(a.get('baseline_points',3)),a.get('baseline_gain_tolerance_db',.2)))
        result['p1db_status']=result.pop('status')
        values,status=at_pout(group,a['target_pout_dbm'])
        result.update(target_pout_dbm=a['target_pout_dbm'],target_status=status)
        if values:
            result.update(values)
        required= a.get('min_p1db_output_dbm') is not None and a.get('aclr_limit_dbc') is not None
        result['assessment']='not_evaluable'
        if values:
            result['assessment']='provisional_missing_criteria' if not required else 'failed'
            if required and result['p1db_status']=='found' and result['p1db_output_dbm']>=a['min_p1db_output_dbm'] and max(values['aclr_lower_dbc'],values['aclr_upper_dbc'])<=a['aclr_limit_dbc']:
                result['assessment']='qualified'
        results.append(result)
    eligible=sorted((r for r in results if r['assessment']=='qualified'),key=lambda r:r['total_idd_a'])
    for rank,r in enumerate(eligible,1):
        r['qualified_current_rank']=rank
    return results


def excel_report(rows,results,c,path):
    from openpyxl import Workbook
    from openpyxl.chart import ScatterChart, Series, Reference
    wb=Workbook();wb.remove(wb.active)
    meta=wb.create_sheet('說明與設定')
    for pair in [('模式',c['mode']),('資料聲明','simulation 為任意示意，非實測；hardware 為使用者命令讀值'),
                 ('RF 設定','rf_set_dbm 是命令設定值；Pin/Pout 必須是讀回值'),
                 ('3 A 策略','僅記錄事件，不停止；不取代任何硬體保護'),
                 ('P1dB','低功率前 N 點平均增益作基準；僅在 1 dB 下降被夾住時插值'),
                 ('相同 Pout','僅對單調曲線內插，不外推；ACLR 採負 dBc，兩側取較差值'),
                 ('條件',f"frequency_hz={c['frequency_hz']}; waveform={c['waveform']}; vdd={c['vdd_v']}; load={c['load_description']}")]:
        meta.append(pair)
    for name,stage in [('靜態資料','static'),('RF資料','rf')]:
        ws=wb.create_sheet(name);ws.append(FIELDS)
        for r in sorted((r for r in rows if r['stage']==stage),key=lambda r:(r['bias_id'],r['pin_dbm'] or 0)):
            ws.append([r[k] for k in FIELDS])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    headers=['bias_id','main_vgs_v','peaking_vgs_v','p1db_status','small_signal_gain_db','p1db_input_dbm','p1db_output_dbm','target_pout_dbm','target_status','total_idd_a','aclr_lower_dbc','aclr_upper_dbc','assessment','qualified_current_rank']
    ws=wb.create_sheet('工作點比較');ws.append(headers)
    for r in results:ws.append([r.get(k) for k in headers])
    events=wb.create_sheet('電流事件');events.append(FIELDS)
    for r in rows:
        if r['event']:events.append([r[k] for k in FIELDS])
    charts=wb.create_sheet('曲線圖')
    charts.append(['圖表由工作簿內數據連結；模擬值不代表實際 PA。最多顯示前 40 組曲線，全部資料仍保留。'])
    data=wb['RF資料'];groups={}
    for n in range(2,data.max_row+1):groups.setdefault(data.cell(n,4).value,[]).append(n)
    specs=[('增益壓縮曲線族','pin_dbm','gain_db'),('Pin–Pout 曲線族','pin_dbm','pout_dbm'),
           ('相同 Pout 工作電流','pout_dbm','total_idd_a'),('下鄰道 ACLR','pout_dbm','aclr_lower_dbc'),('上鄰道 ACLR','pout_dbm','aclr_upper_dbc')]
    for i,(title,x,y) in enumerate(specs):
        chart=ScatterChart();chart.title=title;chart.x_axis.title=x;chart.y_axis.title=y;chart.width=23;chart.height=12
        for bid,indices in list(groups.items())[:40]:
            start,end=indices[0],indices[-1]
            xv=Reference(data,min_col=FIELDS.index(x)+1,min_row=start,max_row=end)
            yv=Reference(data,min_col=FIELDS.index(y)+1,min_row=start,max_row=end)
            label=f"#{bid} Main={data.cell(start,5).value}, Peak={data.cell(start,6).value} V"
            chart.series.append(Series(yv,xv,title=label))
        charts.add_chart(chart,f'A{3+25*i}')
    static=wb['靜態資料']
    for offset,(x,fixed) in enumerate([('main_vgs_v','peaking_vgs_v'),('peaking_vgs_v','main_vgs_v')]):
        chart=ScatterChart();chart.title=x+'–總 Iddq（固定另一支路）';chart.x_axis.title=x;chart.y_axis.title='total_idd_a';chart.width=23;chart.height=12
        grouped={}
        for r in rows:
            if r['stage']=='static':grouped.setdefault(r[fixed],[]).append(r)
        helper=wb.create_sheet('靜態曲線數據'+str(offset));helper.append([x,'total_idd_a'])
        for value,group in list(grouped.items())[:40]:
            start=helper.max_row+1
            for r in sorted(group,key=lambda r:r[x]):helper.append([r[x],r['total_idd_a']])
            chart.series.append(Series(Reference(helper,min_col=2,min_row=start,max_row=helper.max_row),Reference(helper,min_col=1,min_row=start,max_row=helper.max_row),title=f'{fixed}={value} V'))
        charts.add_chart(chart,f'Z{3+25*offset}')
    for ws in wb:
        ws.freeze_panes='A2'
    path=Path(path)
    if path.exists():raise FileExistsError(path)
    wb.save(path)


def execute(c,folder):
    # Validate all dependencies before touching hardware.
    try:
        import openpyxl  # noqa
    except ImportError as exc:
        raise RuntimeError('缺少 Excel 套件，請執行：python -m pip install -r requirements.txt') from exc
    powers=validate(c)
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=False)
    adapter=CommandAdapter(c) if c['mode']=='hardware' else None
    rows=[];fault=None;cleanup=[]
    raw=folder/'measurements.csv'
    with raw.open('x',newline='',encoding='utf-8-sig') as f:
        writer=csv.DictWriter(f,fieldnames=FIELDS);writer.writeheader()
        def record(r):
            rows.append(r);writer.writerow(r);f.flush()
        try:
            if adapter:adapter.call('rf_off')
            for i,p in enumerate(c['bias_points'],1):
                if adapter:
                    adapter.call('rf_off');adapter.call('set_bias',**p);time.sleep(c['settle_seconds'])
                    dc=adapter.call('read_dc')
                else:dc=sim_dc(c,p)
                record(capture(c,'static',dc,p,i))
                if adapter:
                    adapter.call('set_rf',power=min(powers))
                    adapter.call('rf_on')
                for power in powers:
                    if adapter:
                        adapter.call('set_rf',power=power);time.sleep(c['settle_seconds'])
                        dc=adapter.call('read_dc');rf=adapter.call('read_rf')
                    else:dc=sim_dc(c,p,power);rf=sim_rf(p,power)
                    record(capture(c,'rf',dc,p,i,power,rf))
                if adapter:adapter.call('rf_off')
        except Exception as exc:
            fault=exc
        finally:
            if adapter:
                for key in ('rf_off','safe_bias'):
                    try:adapter.call(key)
                    except Exception as exc:cleanup.append(f'{key}: {type(exc).__name__}')
    (folder/'run_status.json').write_text(json.dumps(dict(completed=fault is None and not cleanup,
          error_type=type(fault).__name__ if fault else None,cleanup_errors=cleanup),indent=2),encoding='utf-8')
    (folder/'config.json').write_text(json.dumps(c,ensure_ascii=False,indent=2),encoding='utf-8')
    results=analyze(rows,c) if fault is None else []
    excel_report(rows,results,c,folder/'report.xlsx')
    if fault:raise RuntimeError('執行失敗；已保留部分資料與報表：'+str(folder)) from fault
    if cleanup:raise RuntimeError('結束命令失敗，請檢查硬體狀態；詳見 run_status.json')
    return folder/'report.xlsx'


def save_workflow_config(content,path):
    """Save draft settings without running commands or requiring live hardware."""
    import os
    import tempfile
    data=json.loads(content)
    if not isinstance(data,dict):
        raise ValueError('設定 JSON 最外層必須為物件（大括號）。')
    target=Path(path)
    temporary=None
    try:
        with tempfile.NamedTemporaryFile(mode='w',encoding='utf-8',
                dir=target.parent,prefix='.'+target.name+'.',suffix='.tmp',delete=False) as f:
            temporary=Path(f.name)
            json.dump(data,f,ensure_ascii=False,indent=2,allow_nan=False)
            f.write('\n')
        os.replace(temporary,target)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return target


def gui():
    import tkinter as tk
    from tkinter import filedialog,messagebox,ttk
    root=tk.Tk();root.title('Doherty 完整量測／Excel 報表')
    text=tk.Text(root,width=100,height=34);text.pack(fill='both',expand=True)
    example=Path(__file__).with_name('workflow.example.json')
    if example.exists():text.insert('1.0',example.read_text(encoding='utf-8'))
    current_path=None
    status=tk.StringVar(value='尚未儲存；修改設定後請按儲存。')
    ttk.Label(root,textvariable=status,wraplength=800).pack(fill='x')
    def load():
        nonlocal current_path
        path=filedialog.askopenfilename(filetypes=[('JSON','*.json')])
        if not path:return
        try:
            content=Path(path).read_text(encoding='utf-8-sig')
            if not isinstance(json.loads(content),dict):
                raise ValueError('設定 JSON 最外層必須為物件。')
            text.delete('1.0','end');text.insert('1.0',content)
            current_path=Path(path)
            status.set('已載入：'+str(current_path))
        except (ValueError,OSError) as exc:
            messagebox.showerror('載入失敗',str(exc))
    def save(as_new=False):
        nonlocal current_path
        try:
            content=text.get('1.0','end')
            if not isinstance(json.loads(content),dict):
                raise ValueError('設定 JSON 最外層必須為物件。')
            path=current_path
            if as_new or path is None:
                selected=filedialog.asksaveasfilename(title='儲存設定 JSON',
                    defaultextension='.json',filetypes=[('JSON','*.json')],
                    initialfile=current_path.name if current_path else 'workflow_settings.json')
                if not selected:return
                path=Path(selected)
            current_path=save_workflow_config(content,path)
            status.set('已儲存：'+str(current_path)+'（未執行量測）')
        except (ValueError,OSError) as exc:
            messagebox.showerror('儲存失敗',str(exc))
    def run():
        try:
            c=json.loads(text.get('1.0','end'));validate(c)
            if not messagebox.askokcancel('確認設定',f"模式：{c['mode']}；偏壓 {len(c['bias_points'])} 組。\n3 A 僅記錄事件、不停止。硬體模式會執行設定的本機命令。"):
                return
            parent=filedialog.askdirectory(title='選報表輸出父資料夾')
            if not parent:return
            path=execute(c,Path(parent)/datetime.now().strftime('pa-%Y%m%d-%H%M%S-%f'))
            messagebox.showinfo('完成',str(path))
        except Exception as exc:messagebox.showerror('失敗',str(exc))
    ttk.Button(root,text='載入設定 JSON',command=load).pack(side='left')
    ttk.Button(root,text='儲存設定 JSON',command=save).pack(side='left')
    ttk.Button(root,text='另存新檔',command=lambda:save(True)).pack(side='left')
    ttk.Button(root,text='執行量測／生成含曲線 Excel',command=run).pack(side='left')
    root.mainloop()


if __name__=='__main__':
    import sys
    if len(sys.argv)==1:gui()
    else:
        parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--config',type=Path,required=True);parser.add_argument('--output-dir',type=Path,required=True)
        a=parser.parse_args()
        try:print(execute(json.loads(a.config.read_text(encoding='utf-8')),a.output_dir))
        except Exception as exc:parser.exit(1,str(exc)+'\n')
